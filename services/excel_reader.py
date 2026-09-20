from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from typing import BinaryIO, Union
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from config import HEADER_ALIASES, REQUIRED_COLUMNS
from utils.formatting import clean_cell, normalize_grade, normalize_header

Source = Union[str, BinaryIO]
HEADER_NAME_VALUES = {"name", "student", "student name"}
HEADER_ROLL_VALUES = {
    "roll no",
    "roll number",
    "roll",
    "roll_no",
    "admission no",
    "reg no",
    "ht no",
}
NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


@dataclass
class RawStudent:
    name: str
    roll_no: str
    grade: str
    sheet: str
    row: int


@dataclass
class SheetExtract:
    name: str
    original_name: str
    is_empty: bool
    missing_columns: list[str] = field(default_factory=list)
    records: list[RawStudent] = field(default_factory=list)
    header_row: int | None = None


@dataclass
class WorkbookExtract:
    sheets: list[SheetExtract]
    sheets_read: list[str]
    sheets_skipped: list[str]


def _bytes(source: Source) -> bytes:
    if hasattr(source, "read"):
        source.seek(0)
        data = source.read()
        source.seek(0)
        return data
    with open(source, "rb") as handle:
        return handle.read()


def _xml_sheet_max_rows(data: bytes) -> dict[str, int]:
    result: dict[str, int] = {}
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            rels_xml = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            rid_to_target = {
                rel.attrib["Id"]: rel.attrib["Target"]
                for rel in rels_xml
                if "Id" in rel.attrib and "Target" in rel.attrib
            }
            sheets = workbook_xml.find("m:sheets", NS)
            if sheets is None:
                return result
            for sheet in sheets:
                name = sheet.attrib.get("name", "")
                rid = sheet.attrib.get(REL_ID, "")
                target = rid_to_target.get(rid, "")
                if not target:
                    continue
                if not target.startswith("xl/"):
                    target = "xl/" + target.lstrip("/")
                xml = archive.read(target)
                row_ids = [int(num) for num in re.findall(br'<row[^>]*\br="(\d+)"', xml)]
                cell_ids = [int(num) for num in re.findall(br'<c[^>]*\br="[A-Z]{1,3}(\d+)"', xml)]
                result[name] = max(row_ids + cell_ids) if (row_ids or cell_ids) else 1
    except Exception:
        return result
    return result


def _header_map(values: list) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for col_idx, value in enumerate(values):
        key = HEADER_ALIASES.get(normalize_header(value))
        if key and key not in mapping:
            mapping[key] = col_idx
    if {"roll_no", "name", "grade"}.issubset(mapping):
        return mapping
    return None


def _is_header_row(values: list, mapping: dict[str, int]) -> bool:
    roll = normalize_header(values[mapping["roll_no"]] if mapping["roll_no"] < len(values) else "")
    name = normalize_header(values[mapping["name"]] if mapping["name"] < len(values) else "")
    return roll in HEADER_ROLL_VALUES and name in HEADER_NAME_VALUES


def _infer_mapping(values: list) -> dict[str, int] | None:
    filled = [idx for idx, value in enumerate(values) if clean_cell(value)]
    if len(filled) < 3:
        return None
    start = filled[0]
    if start + 2 >= len(values):
        return None
    first = clean_cell(values[start])
    if first.isdigit() and int(first) in {0, 1} and start + 3 < len(values):
        maybe_roll = clean_cell(values[start + 1])
        if maybe_roll and not maybe_roll.isdigit():
            return {"roll_no": start + 1, "name": start + 2, "grade": start + 3}
    return {"roll_no": start, "name": start + 1, "grade": start + 2}


def _student_from_values(
    values: list,
    mapping: dict[str, int],
    sheet_name: str,
    excel_row: int,
) -> RawStudent | None:
    def cell(key: str):
        idx = mapping[key]
        return values[idx] if idx < len(values) else None

    name_val = clean_cell(cell("name"))
    roll_val = clean_cell(cell("roll_no"))
    sheet_grade = normalize_grade(sheet_name)
    if not sheet_grade.isdigit():
        sheet_grade = ""
    grade_val = normalize_grade(cell("grade")) or sheet_grade
    if not name_val and not roll_val:
        return None
    if _is_header_row(values, mapping):
        return None
    return RawStudent(
        name=name_val,
        roll_no=roll_val,
        grade=grade_val,
        sheet=sheet_name,
        row=excel_row,
    )


def _sheet_last_row(sheet: Worksheet, xml_last: int) -> int:
    last = max(sheet.max_row or 1, xml_last, 1)
    for table in getattr(sheet, "tables", {}).values():
        try:
            _min_col, _min_row, _max_col, max_row = range_boundaries(table.ref)
            last = max(last, max_row)
        except Exception:
            continue
    return last


def _extract_openpyxl_sheet(sheet: Worksheet, xml_last: int = 1) -> SheetExtract:
    original_name = sheet.title
    name = original_name.strip()
    max_col = max(sheet.max_column or 1, 12)
    last_row = _sheet_last_row(sheet, xml_last)
    rows: list[tuple[int, list]] = []
    empty_run = 0
    for row_idx in range(1, last_row + 1):
        values = [sheet.cell(row_idx, col).value for col in range(1, max_col + 1)]
        if not any(clean_cell(value) for value in values):
            empty_run += 1
            continue
        empty_run = 0
        rows.append((row_idx, values))

    if not rows:
        return SheetExtract(name=name, original_name=original_name, is_empty=True)

    mapping = None
    header_row = None
    for excel_row, values in rows[:25]:
        found = _header_map(values)
        if found and _is_header_row(values, found):
            mapping = found
            header_row = excel_row
            break
    if mapping is None:
        mapping = _infer_mapping(rows[0][1])
    if mapping is None:
        return SheetExtract(
            name=name,
            original_name=original_name,
            is_empty=False,
            missing_columns=list(REQUIRED_COLUMNS),
        )

    records = []
    for excel_row, values in rows:
        student = _student_from_values(values, mapping, name, excel_row)
        if student:
            records.append(student)
    return SheetExtract(
        name=name,
        original_name=original_name,
        is_empty=False,
        records=records,
        header_row=header_row,
    )


def _extract_pandas_sheet(original_name: str, frame: pd.DataFrame) -> list[RawStudent]:
    name = original_name.strip()
    if frame is None or frame.empty or frame.dropna(how="all").empty:
        return []
    df = frame.where(pd.notna(frame), None)
    mapping = None
    for idx in range(min(len(df), 25)):
        values = [df.iat[idx, col] for col in range(df.shape[1])]
        found = _header_map(values)
        if found and _is_header_row(values, found):
            mapping = found
            break
    if mapping is None:
        mapping = _infer_mapping([df.iat[0, col] for col in range(df.shape[1])])
    if mapping is None:
        return []
    records = []
    for idx in range(len(df)):
        values = [df.iat[idx, col] for col in range(df.shape[1])]
        student = _student_from_values(values, mapping, name, idx + 1)
        if student:
            records.append(student)
    return records


def _merge_records(primary: list[RawStudent], extra: list[RawStudent]) -> list[RawStudent]:
    seen = {(item.grade, item.roll_no, item.name) for item in primary}
    merged = list(primary)
    for item in extra:
        key = (item.grade, item.roll_no, item.name)
        if key in seen:
            continue
        seen.add(key)
        merged.append(item)
    return merged


def read_workbook(source: Source) -> WorkbookExtract:
    data = _bytes(source)
    xml_max_rows = _xml_sheet_max_rows(data)
    workbook = load_workbook(BytesIO(data), data_only=False)
    try:
        pandas_sheets = pd.read_excel(
            BytesIO(data),
            sheet_name=None,
            header=None,
            dtype=object,
            engine="openpyxl",
            keep_default_na=False,
        )
        extracts: list[SheetExtract] = []
        sheets_read: list[str] = []
        sheets_skipped: list[str] = []
        for sheet in workbook.worksheets:
            extract = _extract_openpyxl_sheet(sheet, xml_max_rows.get(sheet.title, 1))
            pandas_records = _extract_pandas_sheet(
                sheet.title, pandas_sheets.get(sheet.title)
            )
            if extract.is_empty and not pandas_records:
                sheets_skipped.append(extract.name or extract.original_name)
                continue
            if extract.missing_columns and not pandas_records:
                extracts.append(extract)
                sheets_read.append(extract.name)
                continue
            extract.records = _merge_records(extract.records, pandas_records)
            extract.is_empty = False
            extract.missing_columns = []
            extracts.append(extract)
            sheets_read.append(extract.name)
        return WorkbookExtract(
            sheets=extracts,
            sheets_read=sheets_read,
            sheets_skipped=sheets_skipped,
        )
    finally:
        workbook.close()
