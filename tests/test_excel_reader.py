from utils.formatting import normalize_grade

from services.excel_reader import read_workbook


def test_reads_all_sheets_and_strips_sheet_names(tmp_workbook):
    extract = read_workbook(tmp_workbook())
    assert "2nd" in extract.sheets_read
    assert "3rd" in extract.sheets_read
    assert "1st" in extract.sheets_read
    assert "Empty" in extract.sheets_skipped


def test_ignores_empty_rows_and_preserves_roll_text(tmp_workbook):
    extract = read_workbook(tmp_workbook())
    class3 = next(sheet for sheet in extract.sheets if sheet.name == "3rd")
    rolls = [item.roll_no for item in class3.records]
    assert "JPS260301" in rolls
    assert all(item.grade == "3" for item in class3.records)


def test_missing_required_columns(tmp_workbook):
    extract = read_workbook(tmp_workbook("missing_columns"))
    bad = next(sheet for sheet in extract.sheets if sheet.name == "Bad")
    assert bad.missing_columns


def test_counts_every_student_when_header_exists(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "with_header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1st"
    sheet.append(["Roll No", "Name", "Grade"])
    sheet.append(["JPS260101", "ANUSHA R", 1])
    sheet.append(["JPS260102", "BRIAN K", 1])
    sheet.append(["JPS260103", "CHARVI M", 1])
    workbook.save(path)
    extract = read_workbook(path)
    assert len(extract.sheets[0].records) == 3


def test_counts_every_student_when_header_is_missing(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "no_header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1st"
    sheet.append(["JPS260101", "ANUSHA R", 1])
    sheet.append(["JPS260102", "BRIAN K", 1])
    sheet.append(["JPS260103", "CHARVI M", 1])
    workbook.save(path)
    extract = read_workbook(path)
    assert len(extract.sheets[0].records) == 3
    assert extract.sheets[0].records[0].name == "ANUSHA R"


def test_normalizes_ordinal_grades(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "ordinal.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1st"
    sheet.append(["Roll No", "Name", "Grade"])
    sheet.append(["JPS260101", "ANUSHA R", "1st"])
    sheet.append(["JPS260102", "BRIAN K", 1])
    workbook.save(path)
    extract = read_workbook(path)
    grades = {item.grade for item in extract.sheets[0].records}
    assert grades == {"1"}
    assert len(extract.sheets[0].records) == 2


def test_counts_104_class_five_students(tmp_path):
    from openpyxl import Workbook
    from services.student_processor import group_students
    from services.validator import validate_workbook

    path = tmp_path / "class5.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "5th"
    sheet.append(["Roll No", "Name", "Grade"])
    for index in range(1, 105):
        sheet.append([f"JPS2605{index:03d}", f"STUDENT {index}", 5])
    workbook.save(path)
    result = validate_workbook(read_workbook(path))
    grouped = group_students(result.students)
    assert result.ok
    assert len(grouped["5"]) == 104


def test_normalize_grade_values():
    assert normalize_grade("1st") == "1"
    assert normalize_grade("2nd") == "2"
    assert normalize_grade("3rd") == "3"
    assert normalize_grade("Grade 6") == "6"
    assert normalize_grade(1) == "1"
