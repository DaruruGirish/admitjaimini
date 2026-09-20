from pathlib import Path

from openpyxl import Workbook


def build_sample_workbook(path: str | Path, include_issues: str | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    class3 = workbook.active
    class3.title = "3rd"
    class3.append(["Roll No", "Name", "Grade"])
    class3_rows = [
        ("JPS260305", "ANVITHA B", 3),
        ("JPS260301", "AMITH A", 3),
        ("JPS260310", "RITHVIK HENDRY S", 3),
        ("JPS260302", "AMOGH C N", 3),
        ("JPS260315", "THANVITHA GOWDA G", 3),
        ("JPS260304", "SHARATH V", 3),
        ("JPS260318", "KARTHIKEYA VENKATESH GOWDA", 3),
        ("JPS260307", "SHREYAS N R", 3),
        ("JPS260303", "RISHVITH P", 3),
        ("JPS260320", "YUGANTH D", 3),
        ("JPS260312", "TAMILARASI G K", 3),
        ("JPS260309", "THIPPESWAMY T", 3),
    ]
    for row in class3_rows:
        class3.append(list(row))

    class2 = workbook.create_sheet("2nd ")
    class2.append(["Roll No", "Name", "Grade"])
    class2_rows = [
        ("JPS260205", "DIYA S", 2),
        ("JPS260201", "AARAV K", 2),
        ("JPS260204", "CHARVI M", 2),
        ("JPS260202", "BHAVYA R", 2),
        ("JPS260203", "CHETHAN P", 2),
    ]
    for row in class2_rows:
        class2.append(list(row))

    class1 = workbook.create_sheet("1st")
    class1.append(["Roll No", "Name", "Grade"])
    class1_rows = [
        ("JPS260103", "NIDHI A", 1),
        ("JPS260101", "ISHAN V", 1),
        ("JPS260102", "MEERA L", 1),
    ]
    for row in class1_rows:
        class1.append(list(row))

    workbook.create_sheet("Empty")

    if include_issues == "missing_name":
        class1.append(["JPS260199", "", 1])
    elif include_issues == "missing_roll":
        class1.append(["", "MISSING ROLL", 1])
    elif include_issues == "missing_grade":
        class1.append(["JPS260198", "MISSING GRADE", ""])
    elif include_issues == "duplicate_roll":
        class3.append(["JPS260301", "DUPLICATE STUDENT", 3])
    elif include_issues == "missing_columns":
        bad = workbook.create_sheet("Bad")
        bad.append(["Student", "Section"])
        bad.append(["TEST", "A"])

    workbook.save(path)
    return path
